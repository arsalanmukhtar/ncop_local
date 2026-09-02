# project/ncop_internal/flood_report_export.py
# ---------------------------------------------------------------------------
# Builds a self-contained, ONE-PAGE Word (.docx) briefing report for one
# finished flood-model job — dark banner header, map snapshot + legend,
# a compact scenario/accuracy table, a dense colour-coded KPI comparison
# table (the SAME buildings/population/roads/bridges/hospitals exposure
# numbers already shown in the results panel), and a short, prioritized,
# data-driven mitigation-measures list. New file, not folded into
# flood_model_views.py — same "new subsystem gets its own file" precedent
# that file's own module docstring already establishes.
#
# Deliberately dense/compact throughout (narrow margins, small fonts,
# tight spacing, one wide KPI table instead of a dozen labeled rows) —
# a one-page field-briefing layout, not a multi-page narrative report.
# Full per-feature attribute detail (schools/settlements/drainage
# subtype breakdown, per-building geometry) is already available via the
# separate GeoJSON export (FloodModelExportView) — this report is
# deliberately the "at a glance" summary of it, not a duplicate.
#
# Reads _JOBS[job_id]["result"] directly (the SAME in-memory payload
# FloodModelStatusView/FloodModelExportView already read) — exports
# exactly what the user already saw, nothing recomputed, nothing that
# could drift from the displayed result.
#
# The map image is supplied by the CALLER (FloodModelReportExportView),
# captured client-side from the live Mapbox canvas — this module only
# decodes/embeds whatever PNG bytes it's given; it never renders a map
# itself. A missing/malformed image degrades the report to skip that one
# section rather than failing the whole export.
# ---------------------------------------------------------------------------

import base64
import logging
import re
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import flood_model

logger = logging.getLogger(__name__)

# Same severity palette flood-model-control.js's own AHP_ZONE_COLORS and
# flood_model.py's own binary-mask color-relief ramp (_binary_mask_ramp_
# file: RGB(30,100,200)) already use — kept here too, not reinvented, so
# the report's swatches/rows match exactly what the map layer itself
# rendered on screen.
_AHP_ZONE_COLORS_HEX = {"low": "FFC300", "medium": "FF5733", "high": "C70039"}
_FLOOD_EXTENT_COLOR_HEX = "1E64C8"

# Briefing palette — dark navy banner/header bars, light steel-grey
# alternating row shading, white text on dark. Not tied to any real
# security classification (this is an internal planning-tool report,
# not classified material) — purely the dense/dark/tabular VISUAL style
# requested, applied honestly on top of real computed data.
_BANNER_BG = "1B2838"
_HEADER_ROW_BG = "2C3E50"
_ALT_ROW_BG = "EDEFF2"
_WHITE = "FFFFFF"

_MODE_LABELS = {
    "fixed_threshold": "Fixed HAND threshold (susceptibility)",
    "discharge_driven": "Rainfall scenario (discharge-driven)",
    "riverine": "Riverine (live gauge)",
    "ahp_susceptibility": "AHP susceptibility (literature-weighted)",
}

_ZONE_ORDER = ["low", "medium", "high"]
_ZONE_LABELS = {"low": "LOW", "medium": "MED", "high": "HIGH"}

# KPI columns shown in the dense comparison table — the ~8 most
# operationally relevant numbers, not the full attribute set (that's
# what the GeoJSON export is for — see module docstring).
_KPI_COLUMNS = [
    ("area_km2", "Area\n(km²)"),
    ("population", "Pop."),
    ("children", "Child\n<5"),
    ("elderly", "Eld.\n65+"),
    ("buildings", "Bldgs"),
    ("roads_km", "Roads\n(km)"),
    ("bridges", "Bridges"),
    ("hospitals", "Hosp."),
]


# ---------------------------------------------------------------------------
# Small python-docx helpers
# ---------------------------------------------------------------------------

def _set_cell_shading(cell, hex_color):
    """python-docx has no built-in cell-background API — shading is set
    via a raw OOXML <w:shd> element appended to the cell's own tcPr, the
    standard, widely-documented workaround (not a hack specific to this
    app)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, *, bold=False, color_hex=None, size=8, align=None):
    """Replaces a cell's default empty paragraph/run with one styled
    run — avoids python-docx's own `cell.text = ...` shortcut, which
    resets any run-level formatting already applied, making it awkward
    to combine with shading/font-size on the same cell."""
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run("" if text is None else str(text))
    run.font.size = Pt(size)
    run.bold = bold
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)
    return run


def _fmt_num(value):
    """Compact numeric formatting for the dense KPI table — thousands
    separators, no decimals for whole-ish numbers (a KPI strip has no
    room for ".00"), "—" for missing."""
    if value is None or value == "":
        return "—"
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    if f == int(f):
        return f"{int(f):,}"
    return f"{f:,.1f}"


def _fmt(value, suffix=""):
    if value is None or value == "":
        return "—"
    if isinstance(value, float):
        return f"{value:,.2f}{suffix}"
    if isinstance(value, int):
        return f"{value:,}{suffix}"
    return f"{value}{suffix}"


def _in_zone_count(block):
    """block is one of osm_infrastructure["bridges"/"hospitals"] or
    infrastructure["airports"/"schools"/"settlements"] — both share the
    {"total_in_aoi", "in_flood_zone": [...]} shape (fetch_existing_
    infrastructure_exposure / the bridges+hospitals Overpass fetch in
    flood_exposure.py both build it identically). Returns None (not 0)
    when the block itself is missing/empty, so the caller can tell
    "genuinely zero" apart from "no data fetched for this run"."""
    if not isinstance(block, dict) or "in_flood_zone" not in block:
        return None
    return len(block["in_flood_zone"] or [])


def _decode_map_image(data_url):
    """data_url: a 'data:image/png;base64,...' string captured client-
    side from the live Mapbox canvas (flood-model-control.js's own
    #captureMapSnapshot). Returns raw image bytes, or None for a
    missing/malformed value."""
    if not data_url or not isinstance(data_url, str):
        return None
    m = re.match(r"^data:image/(png|jpeg);base64,(.+)$", data_url, re.DOTALL)
    if not m:
        return None
    try:
        return base64.b64decode(m.group(2))
    except Exception:
        logger.warning("flood_report_export: could not decode map_image data URL", exc_info=True)
        return None


# ---------------------------------------------------------------------------
# Exposure -> KPI values / mitigation measures (shared by the catchment-
# wide report and, for AHP mode, each individual zone — both are fed the
# exact same underlying exposure-block shape).
# ---------------------------------------------------------------------------

def _kpi_values(exposure):
    """Flattens one exposure-report dict into the compact numeric set
    the KPI table renders — a small, fixed set of the most
    operationally relevant numbers (see _KPI_COLUMNS), not the full
    attribute breakdown."""
    exposure = exposure or {}
    buildings = exposure.get("buildings") or {}
    population = exposure.get("population") or {}
    roads = exposure.get("roads") or {}
    osm = exposure.get("osm_infrastructure") or {}
    return {
        "area_km2": exposure.get("flood_zone_km2"),
        "population": population.get("total"),
        "children": population.get("under5"),
        "elderly": population.get("elderly"),
        "buildings": buildings.get("buildings_in_flood_zone"),
        "roads_km": roads.get("length_in_flood_zone_km"),
        "bridges": _in_zone_count(osm.get("bridges")),
        "hospitals": _in_zone_count(osm.get("hospitals")),
    }


def _mitigation_measures(exposure, max_items=4):
    """Real, standard disaster-risk-reduction practice, prioritized by
    what THIS run's own computed exposure numbers show is actually at
    stake — not a static paragraph unrelated to the result. Every
    measure is well-established, generic DRR/NDMA-style guidance
    (early-warning dissemination, medical-facility continuity,
    alternate-route planning) applied to this catchment's own real
    numbers — no fabricated citations or invented studies. Capped at
    `max_items` and phrased tightly (one line each) to fit a one-page
    briefing layout — the full, unabridged reasoning per measure isn't
    the point of this format; the priority order and the real trigger
    numbers are."""
    exposure = exposure or {}
    population = exposure.get("population") or {}
    roads = exposure.get("roads") or {}
    osm = exposure.get("osm_infrastructure") or {}
    buildings = exposure.get("buildings") or {}
    bridges_n = _in_zone_count(osm.get("bridges"))
    hospitals_n = _in_zone_count(osm.get("hospitals"))

    measures = []
    pop_total = population.get("total")
    if pop_total:
        measures.append(f"~{pop_total:,.0f} people in zone — activate early-warning dissemination & pre-identified evacuation routes.")
    under5, elderly = population.get("under5"), population.get("elderly")
    if under5 or elderly:
        measures.append(f"{under5 or 0:,.0f} children (<5) + {elderly or 0:,.0f} elderly (65+) — priority-assisted evacuation.")
    if hospitals_n:
        measures.append(f"{hospitals_n} hospital(s)/clinic(s) in zone — confirm continuity plan (power, transfer route).")
    if bridges_n:
        measures.append(f"{bridges_n} bridge(s) in zone — pre-monsoon structural/scour inspection; identify alternate crossing.")
    roads_km = roads.get("length_in_flood_zone_km")
    if roads_km:
        measures.append(f"~{roads_km:,.1f} km road network in zone — sign alternate routes in advance.")
    buildings_n = buildings.get("buildings_in_flood_zone")
    if buildings_n:
        measures.append(f"~{buildings_n:,.0f} buildings in zone — ground-floor flood-proofing where feasible.")

    if not measures:
        measures.append("No population/infrastructure exposure data computed for this result.")
    measures = measures[:max_items]
    measures.append(
        "Baseline: file this zonation with local DM focal points; re-run ahead of forecasted heavy rainfall; "
        "terrain-based hazard indication only — not a real-time warning substitute."
    )
    return measures


def _scenario_rows(job, result, mode):
    if mode == "fixed_threshold":
        return [("HAND threshold", _fmt(job.get("threshold_m"), " m"))]

    flood_zone = result.get("flood_zone") or {}
    scenario = flood_zone.get("scenario") or {}

    if mode == "discharge_driven":
        discharge_scenario = job.get("discharge_scenario") or {}
        return [
            ("Rainfall", _fmt(scenario.get("rainfall_mm"), " mm")),
            ("Duration", _fmt(scenario.get("duration_hr"), " hr")),
            ("Source", discharge_scenario.get("rainfall_source", "manual")),
            ("Composite CN", _fmt(scenario.get("composite_cn"))),
        ]

    if mode == "riverine":
        is_geoglows = scenario.get("gauge_source") == "geoglows_forecast"
        return [
            ("Source", "GeoGLOWS forecast" if is_geoglows else "FFD live gauge"),
            ("Water level", _fmt(scenario.get("gauge_height_m"), " m")),
            ("Status", _fmt(scenario.get("ffd_status"))),
            ("Reading time", _fmt(scenario.get("reading_time"))),
        ]

    if mode == "ahp_susceptibility":
        susceptibility = result.get("susceptibility") or {}
        weights = susceptibility.get("weights") or {}
        top_weights = sorted(weights.items(), key=lambda kv: -(kv[1] or 0))[:3]
        return [
            ("Terrain class", _fmt(susceptibility.get("terrain_class"))),
            ("Top AHP factors", ", ".join(f"{k} ({v})" for k, v in top_weights) or "—"),
        ]

    return []


def _accuracy_rows(result, mode):
    """Real, server-computed validation numbers only — never a fabricated
    confidence figure. Pulls from whichever of this result's own fields
    actually carry one (susceptibility.custom_aoi_accuracy for a custom
    AOI's live AUC ensemble check, shape_comparison for the existing-
    hazard-layer overlap check)."""
    rows = []
    if mode == "ahp_susceptibility":
        acc = (result.get("susceptibility") or {}).get("custom_aoi_accuracy")
        if isinstance(acc, dict):
            rows.append(("AHP AUC", _fmt(acc.get("ahp_auc"))))
            rows.append(("Ensemble AUC", _fmt(acc.get("ensemble_auc"))))
    shape_cmp = result.get("shape_comparison")
    if isinstance(shape_cmp, dict) and not shape_cmp.get("error"):
        rows.append(("Overlap vs. hazard layers", _fmt(shape_cmp.get("iou"))))
    return rows


# ---------------------------------------------------------------------------
# Layout builders
# ---------------------------------------------------------------------------

def _configure_page(doc):
    """Narrow margins + a compact base font — the whole point of a one-
    page briefing layout. Word's own default Normal style (11pt, ~10pt
    paragraph spacing) alone would blow the page budget several times
    over for this much content."""
    section = doc.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)


def _no_borders(table):
    """Suppresses the default single-pixel table borders python-docx's
    'Table Grid' style draws — used for pure LAYOUT tables (e.g. the
    map | scenario two-column split) that shouldn't visually read as a
    data table."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)


def _add_banner(doc, catchment_label, flood_type, mode, generated_at):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        _set_cell_shading(row.cells[0], _BANNER_BG)
    _set_cell_text(
        table.rows[0].cells[0], "FLASH-FLOOD EARLY WARNING — BRIEFING REPORT",
        bold=True, color_hex=_WHITE, size=14,
    )
    _set_cell_text(
        table.rows[1].cells[0],
        f"{catchment_label}   |   {flood_type.title()} — {_MODE_LABELS.get(mode, mode or 'Unknown')}   |   "
        f"Generated {generated_at}",
        bold=False, color_hex=_WHITE, size=9,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_map_and_scenario_row(doc, image_bytes, legend_items, scenario_rows, accuracy_rows):
    outer = doc.add_table(rows=1, cols=2)
    outer.autofit = False
    _no_borders(outer)
    left, right = outer.rows[0].cells
    left.width = Inches(2.9)
    right.width = Inches(4.2)

    # --- Left: map thumbnail + legend chips ---
    if image_bytes:
        p = left.paragraphs[0]
        run = p.add_run()
        try:
            run.add_picture(BytesIO(image_bytes), width=Inches(2.75))
        except Exception:
            logger.warning("flood_report_export: could not embed map snapshot image", exc_info=True)
        legend_p = left.add_paragraph()
        legend_p.paragraph_format.space_before = Pt(2)
        for label, hex_color in legend_items:
            swatch_run = legend_p.add_run("  ■ ")
            swatch_run.font.color.rgb = RGBColor.from_string(hex_color)
            swatch_run.font.size = Pt(8)
            label_run = legend_p.add_run(f"{label}   ")
            label_run.font.size = Pt(7.5)
    else:
        p = left.paragraphs[0]
        run = p.add_run("(No map snapshot captured for this export.)")
        run.italic = True
        run.font.size = Pt(8)

    # --- Right: compact Scenario + Accuracy table ---
    combined_rows = [("SCENARIO", None)] + scenario_rows + [("ACCURACY & VALIDATION", None)] + (
        accuracy_rows or [("No live comparison computed", "")]
    )
    inner = right.add_table(rows=0, cols=2)
    inner.autofit = True
    for i, (label, value) in enumerate(combined_rows):
        cells = inner.add_row().cells
        if value is None:
            _set_cell_shading(cells[0], _HEADER_ROW_BG)
            _set_cell_shading(cells[1], _HEADER_ROW_BG)
            _set_cell_text(cells[0], label, bold=True, color_hex=_WHITE, size=8)
            cells[1].text = ""
        else:
            if i % 2 == 0:
                _set_cell_shading(cells[0], _ALT_ROW_BG)
                _set_cell_shading(cells[1], _ALT_ROW_BG)
            _set_cell_text(cells[0], label, bold=True, size=8)
            _set_cell_text(cells[1], "—" if value in (None, "") else str(value), size=8)
    right.paragraphs[0].text = ""


def _add_kpi_table(doc, zone_rows):
    """zone_rows: [(zone_key_or_None, zone_color_hex_or_None, kpi_dict), ...]
    One shared table shape for both a single catchment-wide row (non-AHP
    modes) and a 3-row per-zone comparison (AHP mode) — a genuine
    sector-comparison table either way, exactly the "one dense table"
    shape a field briefing uses instead of narrative paragraphs."""
    has_zone_col = any(z is not None for z, _, _ in zone_rows)
    ncols = len(_KPI_COLUMNS) + (1 if has_zone_col else 0)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = table.rows[0].cells
    col_offset = 0
    if has_zone_col:
        _set_cell_shading(header_cells[0], _HEADER_ROW_BG)
        _set_cell_text(header_cells[0], "ZONE", bold=True, color_hex=_WHITE, size=7.5,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
        col_offset = 1
    for i, (_, col_label) in enumerate(_KPI_COLUMNS):
        cell = header_cells[i + col_offset]
        _set_cell_shading(cell, _HEADER_ROW_BG)
        _set_cell_text(cell, col_label, bold=True, color_hex=_WHITE, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_i, (zone_key, zone_color, kpis) in enumerate(zone_rows):
        cells = table.add_row().cells
        row_bg = _ALT_ROW_BG if row_i % 2 == 1 else None
        if has_zone_col:
            _set_cell_shading(cells[0], zone_color or row_bg or _WHITE)
            _set_cell_text(
                cells[0], _ZONE_LABELS.get(zone_key, zone_key or "—"),
                bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER,
                color_hex=_WHITE if zone_color else None,
            )
        for i, (key, _) in enumerate(_KPI_COLUMNS):
            cell = cells[i + col_offset]
            if row_bg and not (has_zone_col and i == 0):
                _set_cell_shading(cell, row_bg)
            _set_cell_text(cell, _fmt_num(kpis.get(key)), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    return table


def _add_mitigation_list(doc, measures):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run("PRIORITY ACTIONS")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(_BANNER_BG)
    for i, measure in enumerate(measures, start=1):
        mp = doc.add_paragraph()
        mp.paragraph_format.space_after = Pt(1)
        mp.paragraph_format.left_indent = Pt(10)
        num_run = mp.add_run(f"{i}. ")
        num_run.bold = True
        num_run.font.size = Pt(8)
        text_run = mp.add_run(measure)
        text_run.font.size = Pt(8)


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def build_flood_model_docx_report(job, map_image_data_url=None):
    """Builds the full one-page briefing report for one finished job
    (job["status"] == "done" — the caller, FloodModelReportExportView,
    is responsible for that check). Returns raw .docx bytes.

    `map_image_data_url` — optional, a 'data:image/png;base64,...'
    string captured client-side from the live map canvas at the moment
    the user clicked Export. Omitted or undecodable -> the report skips
    the map-snapshot section rather than failing outright.

    Page-fit is engineered (narrow margins, 9pt/8pt fonts, one dense
    KPI table instead of a dozen labeled rows, a 4-item capped
    mitigation list) rather than measured against a real Word layout
    engine — this environment has no LibreOffice/Word available to
    render an actual page count. For the realistic content volume this
    produces (one banner, one map+scenario row, one KPI table of up to
    3 rows, 5 mitigation lines, one footer) the estimated content height
    is roughly half of one Letter page's usable height at these
    settings — comfortable margin, not a hairline fit — but treat "one
    page" as engineered-for, not independently rendered and measured."""
    result = job.get("result") or {}
    mode = result.get("mode")
    catchment_key = job.get("catchment")
    cfg = flood_model.PILOT_CATCHMENTS.get(catchment_key, {})
    catchment_label = cfg.get("label", catchment_key)
    flood_type = cfg.get("flood_type", "flash")
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    doc = Document()
    _configure_page(doc)
    _add_banner(doc, catchment_label, flood_type, mode, generated_at)

    if result.get("no_flood_expected"):
        p = doc.add_paragraph()
        run = p.add_run("NO FLOOD RISK CURRENTLY INDICATED")
        run.bold = True
        run.font.size = Pt(11)
        doc.add_paragraph(result.get("message") or "The fetched scenario value was too low to drive the discharge model.")
        doc_bytes = BytesIO()
        doc.save(doc_bytes)
        return doc_bytes.getvalue()

    image_bytes = _decode_map_image(map_image_data_url)
    legend_items = (
        [(_ZONE_LABELS[z] + " susceptibility", _AHP_ZONE_COLORS_HEX[z]) for z in _ZONE_ORDER]
        if mode == "ahp_susceptibility"
        else [("Flood-prone / inundated extent", _FLOOD_EXTENT_COLOR_HEX)]
    )
    _add_map_and_scenario_row(
        doc, image_bytes, legend_items,
        _scenario_rows(job, result, mode), _accuracy_rows(result, mode),
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    kpi_heading = doc.add_paragraph()
    kpi_run = kpi_heading.add_run("EXPOSURE & INFRASTRUCTURE — IN ZONE")
    kpi_run.bold = True
    kpi_run.font.size = Pt(9)
    kpi_run.font.color.rgb = RGBColor.from_string(_BANNER_BG)

    if mode == "ahp_susceptibility":
        by_zone = (result.get("zone_exposure") or {}).get("by_zone", {})
        zone_rows = [
            (zone, _AHP_ZONE_COLORS_HEX[zone], _kpi_values(by_zone[zone]))
            for zone in _ZONE_ORDER if zone in by_zone
        ]
        if zone_rows:
            _add_kpi_table(doc, zone_rows)
            # A combined mitigation list drawn from the HIGHEST-susceptibility
            # zone actually present (worst-case first — the zone most in need
            # of action, not an arbitrary average across all three).
            worst_zone = next((z for z in reversed(_ZONE_ORDER) if z in by_zone), None)
            measures = _mitigation_measures(by_zone.get(worst_zone))
        else:
            doc.add_paragraph("No zone exposure data computed for this run.")
            measures = _mitigation_measures(None)
    else:
        exposure = result.get("exposure")
        if exposure:
            _add_kpi_table(doc, [(None, None, _kpi_values(exposure))])
        else:
            doc.add_paragraph("No exposure report computed for this run.")
        measures = _mitigation_measures(exposure)

    _add_mitigation_list(doc, measures)

    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(6)
    footer_run = footer_p.add_run(
        "Terrain/hazard-model output for disaster-management planning use — not a certified engineering "
        "flood forecast. Full attribute detail (schools, settlements, drainage breakdown, per-feature "
        "geometry) available via the separate GeoJSON export of this same result."
    )
    footer_run.italic = True
    footer_run.font.size = Pt(6.5)
    footer_run.font.color.rgb = RGBColor.from_string("666666")

    doc_bytes = BytesIO()
    doc.save(doc_bytes)
    return doc_bytes.getvalue()
